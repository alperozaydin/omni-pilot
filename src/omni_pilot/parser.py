"""Parse MacroFactor xlsx exports."""
from __future__ import annotations

import logging
import os
from collections import defaultdict
from collections.abc import Iterable
from datetime import date, datetime
from typing import TypedDict

import openpyxl
import yaml

logger = logging.getLogger(__name__)


class FoodEntry(TypedDict):
    date: str
    time: str
    food_name: str
    serving_size: str
    serving_qty: float
    serving_weight_g: float
    total_weight_g: float
    calories_kcal: float
    protein_g: float
    fat_g: float
    carbs_g: float
    # MacroFactor's own fiber figure for this entry; None when the cell is empty.
    fiber_g: float | None
    # total_weight_g is a placeholder, not what was eaten (see is_placeholder_weight).
    weight_is_placeholder: bool


class LastDayCheck(TypedDict):
    date: str
    logged_kcal: float
    # The day's MacroFactor expenditure (TDEE); None when the export has none for it.
    expenditure_kcal: float | None
    # Logged under PARTIAL_DAY_SHARE of the expenditure: the day was still being logged.
    partial: bool


# A last day logged under this share of its expenditure is taken as still in progress.
# Full days in the exports so far reached at least 79%; the partly logged last
# days 47% and 52% (a lunch-time export), which a cutoff of half let through.
PARTIAL_DAY_SHARE = 0.7

# Meals logged from a photo get a serving weight of 100 g whatever their size,
# while their macros are right for the whole meal. Such an entry gives itself
# away when its protein, fat and carbs weigh more than the entry, which no food
# can; the margin absorbs whole-gram rounding.
PLACEHOLDER_SERVING_G = 100.0
PLACEHOLDER_MACRO_MARGIN = 1.05


# Column indices in the "Food Log" sheet (0-indexed)
_COL_DATE = 0
_COL_TIME = 1
_COL_FOOD_NAME = 2
_COL_SERVING_SIZE = 3
_COL_SERVING_QTY = 4
_COL_SERVING_WEIGHT_G = 5
_COL_CALORIES = 6
_COL_FAT = 7
_COL_CARBS = 8
_COL_PROTEIN = 9
_COL_FIBER = 27


def parse_food_log(xlsx_path: str) -> list[FoodEntry]:
    """Parse the 'Food Log' sheet from a MacroFactor export.

    Returns a list of FoodEntry dicts. Skips rows with no food name.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb["Food Log"]

    entries: list[FoodEntry] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        food_name = row[_COL_FOOD_NAME]
        if food_name is None or str(food_name).strip() == "":
            continue

        food_name = str(food_name).strip()
        serving_qty = float(row[_COL_SERVING_QTY] or 0)
        serving_weight_g = float(row[_COL_SERVING_WEIGHT_G] or 0)
        total_weight_g = serving_qty * serving_weight_g
        logged_fiber = row[_COL_FIBER] if len(row) > _COL_FIBER else None

        if food_name == "Quick Add":
            logger.warning(
                "Row %d: 'Quick Add' entry skipped for micro analysis "
                "(no food name to look up).",
                row_idx,
            )

        calories_kcal = float(row[_COL_CALORIES] or 0)
        protein_g = float(row[_COL_PROTEIN] or 0)
        fat_g = float(row[_COL_FAT] or 0)
        carbs_g = float(row[_COL_CARBS] or 0)
        entries.append(FoodEntry(
            date=str(row[_COL_DATE] or ""),
            time=str(row[_COL_TIME] or ""),
            food_name=food_name,
            serving_size=str(row[_COL_SERVING_SIZE] or ""),
            serving_qty=serving_qty,
            serving_weight_g=serving_weight_g,
            total_weight_g=total_weight_g,
            calories_kcal=calories_kcal,
            protein_g=protein_g,
            fat_g=fat_g,
            carbs_g=carbs_g,
            fiber_g=float(logged_fiber) if logged_fiber is not None else None,
            weight_is_placeholder=is_placeholder_weight(
                serving_weight_g, total_weight_g, protein_g, fat_g, carbs_g,
            ),
        ))

    wb.close()
    return entries


def is_placeholder_weight(
    serving_weight_g: float, total_weight_g: float, protein_g: float, fat_g: float, carbs_g: float,
) -> bool:
    """Whether an entry's weight is MacroFactor's 100 g placeholder for a photo-logged meal."""
    return (
        serving_weight_g == PLACEHOLDER_SERVING_G
        and protein_g + fat_g + carbs_g > PLACEHOLDER_MACRO_MARGIN * total_weight_g
    )


def parse_daily_expenditure(xlsx_path: str) -> dict[str, float]:
    """Read each day's calorie expenditure from the "Quick Export" sheet, keyed by ISO date.

    Uses Expenditure (MacroFactor's TDEE estimate), falling back to Target Calories;
    a day with neither is left out. An export without the sheet gives an empty dict.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    try:
        if "Quick Export" not in wb.sheetnames:
            return {}
        rows = wb["Quick Export"].iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows, ())]
        if "Date" not in header:
            return {}
        col_date = header.index("Date")
        cols_kcal = [header.index(name) for name in ("Expenditure", "Target Calories (kcal)") if name in header]

        expenditure: dict[str, float] = {}
        for row in rows:
            day = row[col_date] if len(row) > col_date else None
            if day is None:
                continue
            day_iso = day.date().isoformat() if isinstance(day, datetime) else (
                day.isoformat() if isinstance(day, date) else str(day)[:10]
            )
            for col in cols_kcal:
                if len(row) > col and row[col]:
                    expenditure[day_iso] = float(row[col])
                    break
        return expenditure
    finally:
        wb.close()


def check_last_day(entries: list[FoodEntry], daily_expenditure: dict[str, float]) -> LastDayCheck | None:
    """Check whether the export's last day was logged only in part (BAR-87).

    An export taken mid-day holds a half-logged last day that would pull every
    average down. That day is partial when its logged calories (every entry,
    Quick Add included) stay under PARTIAL_DAY_SHARE of its expenditure. Without
    an expenditure figure the day counts. Returns None when there are fewer than
    two days, since leaving out the only day would leave nothing to analyse.
    """
    logged_kcal: dict[str, float] = defaultdict(float)
    for entry in entries:
        logged_kcal[entry["date"]] += entry["calories_kcal"]
    if len(logged_kcal) < 2:
        return None

    last_day = max(logged_kcal)
    expenditure = daily_expenditure.get(last_day)
    return LastDayCheck(
        date=last_day,
        logged_kcal=logged_kcal[last_day],
        expenditure_kcal=expenditure,
        partial=expenditure is not None and logged_kcal[last_day] < PARTIAL_DAY_SHARE * expenditure,
    )


def extract_unique_foods(entries: list[FoodEntry]) -> list[str]:
    """Extract sorted list of unique food names from parsed entries."""
    return sorted({e["food_name"] for e in entries})


def food_names_in_exports(paths: Iterable[str]) -> set[str]:
    """The unique food names logged in any of the given MacroFactor exports."""
    names: set[str] = set()
    for path in paths:
        names.update(extract_unique_foods(parse_food_log(path)))
    return names


def generate_food_mappings(
    foods: list[str],
    existing_mappings: dict[str, str],
    output_path: str,
) -> bool:
    """Generate or merge food_mappings.yaml.

    Only foods with valid non-empty mappings are saved.
    Returns True if the file was created or updated, False if it was already up to date.
    """
    merged: dict[str, str] = {}
    for food in sorted(foods):
        if food in existing_mappings and existing_mappings[food] and str(existing_mappings[food]).strip():
            merged[food] = str(existing_mappings[food]).strip()

    # Also preserve any existing non-empty mappings not in the current food list
    for food, mapping in existing_mappings.items():
        if food not in merged and mapping and str(mapping).strip():
            merged[food] = str(mapping).strip()

    if os.path.exists(output_path):
        try:
            with open(output_path) as f:
                existing_file_data = yaml.safe_load(f)
            if isinstance(existing_file_data, dict) and existing_file_data.get("mappings") == merged:
                return False
        except Exception:
            pass

    data = {
        "mappings": merged,
    }

    header = (
        "# Written by Omni Pilot: hand edits are overwritten on the next run.\n"
        "# To change a food's USDA lookup text: make update-mapping FOOD=\"...\" QUERY=\"...\"\n"
        "# (without make: python -m omni_pilot.cli update-mapping \"FOOD\" \"QUERY\").\n"
        "# QUERY=skip leaves a food out of the micronutrient analysis.\n\n"
    )

    parent = os.path.dirname(output_path)
    if parent:
        os.makedirs(parent, exist_ok=True)

    with open(output_path, "w") as f:
        f.write(header)
        yaml.dump(data, f, default_flow_style=False, allow_unicode=True, sort_keys=False)

    return True

